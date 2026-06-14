#!/usr/bin/env python3
import argparse
import csv
import json
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import requests
import websocket
from bs4 import BeautifulSoup


DEFAULT_BASE_URL = "https://psa-portal.telekom.de/intranet-ui"


class CdpError(RuntimeError):
    pass


class CdpClient:
    def __init__(self, port, url_contains=None):
        self.port = port
        self.target = self._select_target(url_contains)
        self._next_id = 1
        self.ws = websocket.create_connection(
            self.target["webSocketDebuggerUrl"],
            timeout=30,
            suppress_origin=True,
        )

    def close(self):
        self.ws.close()

    def _select_target(self, url_contains):
        response = requests.get(f"http://127.0.0.1:{self.port}/json/list", timeout=10)
        response.raise_for_status()
        targets = [t for t in response.json() if t.get("type") == "page"]
        if not targets:
            raise CdpError(f"No page targets found on DevTools port {self.port}")
        if url_contains:
            for target in targets:
                if url_contains in target.get("url", ""):
                    return target
        return targets[0]

    def call(self, method, params=None):
        message_id = self._next_id
        self._next_id += 1
        self.ws.send(json.dumps({"id": message_id, "method": method, "params": params or {}}))
        while True:
            raw = self.ws.recv()
            message = json.loads(raw)
            if message.get("id") != message_id:
                continue
            if "error" in message:
                raise CdpError(f"{method} failed: {message['error']}")
            return message.get("result", {})

    def evaluate(self, expression, await_promise=False):
        result = self.call(
            "Runtime.evaluate",
            {
                "expression": expression,
                "awaitPromise": await_promise,
                "returnByValue": True,
            },
        )
        payload = result.get("result", {})
        if "exceptionDetails" in result:
            raise CdpError(json.dumps(result["exceptionDetails"], indent=2))
        return payload.get("value")

    def evaluate_json(self, expression, await_promise=False):
        if await_promise:
            value = self.evaluate(expression, await_promise=True)
            if isinstance(value, str):
                return json.loads(value)
            return value
        value = self.evaluate(f"JSON.stringify(({expression}))", await_promise=await_promise)
        return json.loads(value)


def system_url(base_url, system_id):
    return f"{base_url.rstrip('/')}/my/systems/details/details.xhtml?system={system_id}"


def wait_ready(cdp, expected_url=None, timeout=45):
    deadline = time.time() + timeout
    last_state = None
    while time.time() < deadline:
        state = cdp.evaluate_json(
            "(() => ({href: location.href, title: document.title, readyState: document.readyState}))()"
        )
        last_state = state
        href_ok = expected_url is None or state["href"].startswith(expected_url)
        if href_ok and state["readyState"] == "complete":
            return state
        time.sleep(0.5)
    raise CdpError(f"Page did not become ready in time. Last state: {last_state}")


def navigate(cdp, url):
    cdp.call("Page.navigate", {"url": url})
    return wait_ready(cdp, expected_url=url)


def parse_sheet_ids(href):
    parsed = urlparse(href)
    query = parse_qs(parsed.query)
    return {
        "sheet": query.get("sheet", [""])[0],
        "system": query.get("system", [""])[0],
    }


def snapshot_page(cdp):
    data = cdp.evaluate_json(
        """(() => {
          const links = Array.from(document.links).map((a) => ({
            text: (a.innerText || a.textContent || "").trim(),
            title: a.title || "",
            href: a.href,
            id: a.id || ""
          }));
          const socSheets = links
            .filter((link) => link.href.includes("/socsheet/socsheet.xhtml"))
            .map((link) => ({name: link.text, href: link.href}));
          const text = document.body ? document.body.innerText : "";
          return {
            href: location.href,
            title: document.title,
            readyState: document.readyState,
            capturedAt: new Date().toISOString(),
            noSocs: text.includes("Keine SoCs"),
            text,
            html: document.documentElement ? document.documentElement.outerHTML : "",
            links,
            socSheets
          };
        })()"""
    )
    for sheet in data["socSheets"]:
        sheet.update(parse_sheet_ids(sheet["href"]))
    return data


def write_snapshot(snapshot, out_dir, slug):
    out_dir.mkdir(parents=True, exist_ok=True)
    snapshot_path = out_dir / f"{slug}_page_snapshot.json"
    text_path = out_dir / f"{slug}_page_text.txt"
    manifest_json = out_dir / f"{slug}_soc_manifest.json"
    manifest_tsv = out_dir / f"{slug}_soc_manifest.tsv"

    snapshot_path.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False), encoding="utf-8")
    text_path.write_text(snapshot["text"], encoding="utf-8")
    manifest_json.write_text(
        json.dumps(snapshot["socSheets"], indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    with manifest_tsv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["name", "sheet", "system", "href"], delimiter="\t")
        writer.writeheader()
        for row in snapshot["socSheets"]:
            writer.writerow({key: row.get(key, "") for key in writer.fieldnames})
    return {
        "snapshot": str(snapshot_path),
        "text": str(text_path),
        "manifestJson": str(manifest_json),
        "manifestTsv": str(manifest_tsv),
        "socCount": len(snapshot["socSheets"]),
        "noSocs": snapshot["noSocs"],
    }


def select_all_socs(cdp):
    return cdp.evaluate_json(
        """(() => new Promise((resolve) => {
          const master = document.querySelector('input[type="checkbox"][id*="bulk-toggle-master_input"]');
          const exportButton = Array.from(document.querySelectorAll("button"))
            .find((button) => (button.id || "").includes("bulk-export"));
          const text = document.body ? document.body.innerText : "";
          if (!master || !exportButton) {
            resolve({
              ok: false,
              reason: text.includes("Keine SoCs") ? "no socs" : "bulk controls not found",
              total: 0,
              checked: 0,
              noSocs: text.includes("Keine SoCs")
            });
            return;
          }
          if (!master.checked) {
            master.click();
          }
          setTimeout(() => {
            const boxes = Array.from(document.querySelectorAll('input[type="checkbox"][id*="id_socs_table"]'))
              .filter((box) => !(box.id || "").includes("bulk-toggle-master_input"));
            resolve({
              ok: true,
              total: boxes.length,
              checked: boxes.filter((box) => box.checked).length,
              exportDisabled: exportButton.disabled || exportButton.className.includes("ui-state-disabled"),
              exportButtonId: exportButton.id
            });
          }, 2500);
        }))()""",
        await_promise=True,
    )


def click_bulk_export(cdp):
    return cdp.evaluate_json(
        """(() => new Promise((resolve) => {
          const exportButton = Array.from(document.querySelectorAll("button"))
            .find((button) => (button.id || "").includes("bulk-export"));
          if (!exportButton) {
            resolve({ok: false, reason: "bulk export button not found"});
            return;
          }
          exportButton.click();
          setTimeout(() => resolve({ok: true, clicked: true}), 1000);
        }))()""",
        await_promise=True,
    )


def wait_for_download(download_dir, start_time, timeout=120):
    deadline = time.time() + timeout
    seen_candidate = None
    while time.time() < deadline:
        partials = list(download_dir.glob("*.crdownload"))
        files = [
            path
            for path in download_dir.iterdir()
            if path.is_file()
            and not path.name.endswith(".crdownload")
            and path.stat().st_mtime >= start_time
        ]
        zip_files = [path for path in files if path.suffix.lower() == ".zip"]
        candidates = zip_files or files
        if candidates and not partials:
            newest = max(candidates, key=lambda path: path.stat().st_mtime)
            size = newest.stat().st_size
            time.sleep(1)
            if newest.exists() and newest.stat().st_size == size:
                return newest
            seen_candidate = str(newest)
        time.sleep(1)
    raise CdpError(f"Timed out waiting for download. Last candidate: {seen_candidate}")


def safe_filename(value):
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    value = value.strip("._")
    return value[:80] or "sheet"


def md_escape(value):
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return text


def one_line(value):
    return re.sub(r"\s+", " ", md_escape(value))


def extract_xlsx(xlsx_path, out_dir, slug):
    from openpyxl import load_workbook

    def cell_value(value):
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        if hasattr(value, "isoformat"):
            return value.isoformat()
        return str(value)

    out_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(xlsx_path, data_only=False, read_only=True)
    extracted = {
        "source": str(xlsx_path),
        "extractedAt": datetime.now(timezone.utc).isoformat(),
        "sheets": [],
    }
    for sheet in workbook.worksheets:
        rows = []
        for raw_row in sheet.iter_rows(values_only=True):
            row = [cell_value(value) for value in raw_row]
            while row and row[-1] is None:
                row.pop()
            if any(cell is not None for cell in row):
                rows.append(row)

        sheet_file = out_dir / f"{slug}_{safe_filename(sheet.title)}.csv"
        with sheet_file.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows(rows)

        extracted["sheets"].append(
            {
                "title": sheet.title,
                "csv": str(sheet_file),
                "rowCount": len(rows),
                "rows": rows,
            }
        )

    json_file = out_dir / f"{slug}_workbook.json"
    json_file.write_text(json.dumps(extracted, indent=2, ensure_ascii=False), encoding="utf-8")
    return {
        "source": str(xlsx_path),
        "workbookJson": str(json_file),
        "sheets": [
            {"title": sheet["title"], "csv": sheet["csv"], "rowCount": sheet["rowCount"]}
            for sheet in extracted["sheets"]
        ],
    }


def workbook_requirement_sheets(xlsx_path):
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, data_only=False, read_only=True)
    for sheet in workbook.worksheets:
        if not re.match(r"^\d+\)", sheet.title):
            continue
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 3:
            continue
        yield sheet.title, rows


def requirement_rows(rows):
    parsed = []
    for index, row in enumerate(rows[2:], start=3):
        key = row[0] if len(row) > 0 else None
        text = row[1] if len(row) > 1 else None
        if not key or not text:
            continue
        statement = ""
        statement_columns = [
            ("SOC", "Fulfilled (SoC)", 4),
            ("SOPC", "Partially fulfilled (SoPC)", 5),
            ("SONC", "Not fulfilled (SoNC)", 6),
            ("NA", "Not applicable", 7),
        ]
        for code, label, column in statement_columns:
            value = row[column] if len(row) > column else None
            if isinstance(value, str) and value.strip().lower() == "x":
                statement = label
                break
        parsed.append(
            {
                "excelRow": index,
                "key": md_escape(key),
                "requirement": md_escape(text),
                "statement": statement,
                "explanation": md_escape(row[8] if len(row) > 8 else ""),
                "securityTest": md_escape(row[9] if len(row) > 9 else ""),
                "reviewResult": md_escape(row[12] if len(row) > 12 else ""),
                "reviewExplanation": md_escape(row[13] if len(row) > 13 else ""),
                "requirementId": md_escape(row[14] if len(row) > 14 else ""),
            }
        )
    return parsed


ANSWER_GUIDE = """\
## How To Answer In PSA

Use the live PSA form fields observed through the browser:

- `Statement` is a radio group per requirement with values `SOC`, `SOPC`, `SONC`, and `NA`.
- `SOC` / fulfilled means the requirement is fully implemented for the system scope. The portal does not require a justification for this choice, but keep evidence in the explanation or linked documentation when useful.
- `SOPC` / partially fulfilled means only part of the requirement is implemented. Fill `Begruendung` with what is covered, what is missing, affected scope, compensating controls, and the planned measure/risk handling.
- `SONC` / not fulfilled means the requirement is currently not implemented. Fill `Begruendung` with why, impact, compensating controls, owner, and planned remediation/risk acceptance path.
- `NA` / not applicable means the requirement does not apply to this system context. Fill `Begruendung` with the concrete scoping reason, not just "not used".
- `Notizbereich` is a working note area. It has `Meine Notiz`, a visibility checkbox `Fuer andere sichtbar`, and a read-only field for visible notes from other users. Do not use notes as the authoritative answer.
- Use `Speichern` only after intentional edits. `Speichern und beenden` is disabled until the sheet can be completed.
"""


def requirement_guidance(requirement_text):
    text = requirement_text.lower()
    def has_keyword(keywords):
        short_keywords = {"gke", "iam", "csp", "vpc", "siem", "tls"}
        for keyword in keywords:
            if keyword in short_keywords:
                if re.search(rf"(?<![a-z0-9]){re.escape(keyword)}(?![a-z0-9])", text):
                    return True
            elif keyword in text:
                return True
        return False

    guidance = [
        "Confirm whether the requirement applies to the exact system boundary and operating model.",
        "Collect concrete evidence before selecting the statement value.",
    ]
    if "muss" in text or "müssen" in text:
        guidance.append("This is phrased as mandatory; choose `SOC` only when the control is fully implemented.")
    if "sollte" in text:
        guidance.append("This is phrased as recommended/expected; document the rationale if not implemented exactly as written.")
    keyword_guidance = [
        (("personenbezogen", "datenschutz", "einwilligung", "betroffene"), "Check privacy documentation, data categories, legal basis, retention/deletion, and DSB guidance."),
        (("cloud", "gcp", "csp", "landingzone", "landing zone"), "Distinguish platform/landing-zone responsibility from application/team responsibility in the justification."),
        (("iam", "benutzer", "berechtigung", "authent", "konto", "accounts"), "Use IAM role design, access review, MFA, privileged access, and account lifecycle evidence."),
        (("log", "monitor", "siem", "incident", "vorfall"), "Use monitoring, logging, SIEM forwarding, alert handling, and incident workflow evidence."),
        (("backup", "restore", "wiederherstellung"), "Use backup configuration, restore-test evidence, RPO/RTO, and ownership."),
        (("verschlüssel", "krypto", "tls", "zertifikat", "schlüssel"), "Use crypto policy, TLS/cipher configuration, certificate/key management, and rotation evidence."),
        (("container", "image", "kubernetes", "gke"), "Use cluster policy, image provenance, runtime hardening, admission controls, and patching evidence."),
        (("netz", "firewall", "vpc", "router", "switch", "ip-adresse"), "Use network segmentation, firewall/security group rules, routing, IP management, and exposure evidence."),
        (("cicd", "pipeline", "repository", "deployment"), "Use repository controls, pipeline approvals, artifact provenance, secret handling, and deployment traceability."),
        (("ki-system", "künstliche intelligenz", "machine learning", "modell"), "Clarify whether AI/ML is part of the system function or only supporting operations; justify non-applicability explicitly when relevant."),
    ]
    for keywords, note in keyword_guidance:
        if has_keyword(keywords):
            guidance.append(note)
    guidance.append("If the answer is `SOPC`, `SONC`, or `NA`, write the concrete `Begruendung` in the portal field.")
    return guidance


def render_markdown_from_xlsx(xlsx_path, out_dir, slug):
    out_dir.mkdir(parents=True, exist_ok=True)
    index_lines = [
        f"# {slug} SoC Requirements",
        "",
        f"Source workbook: `{xlsx_path}`",
        "",
        ANSWER_GUIDE.strip(),
        "",
        "## Sheets",
        "",
    ]
    summary = {"source": str(xlsx_path), "generatedAt": datetime.now(timezone.utc).isoformat(), "sheets": []}

    for sheet_title, rows in workbook_requirement_sheets(xlsx_path):
        reqs = requirement_rows(rows)
        sheet_slug = safe_filename(sheet_title)
        sheet_path = out_dir / f"{slug}_{sheet_slug}.md"
        index_lines.append(f"- [{sheet_title}]({sheet_path.name}) - {len(reqs)} requirements")
        lines = [
            f"# {sheet_title}",
            "",
            f"Requirement count: {len(reqs)}",
            "",
            ANSWER_GUIDE.strip(),
            "",
            "## Requirements",
            "",
        ]
        for req in reqs:
            lines.extend(
                [
                    f"### {req['key']}",
                    "",
                    f"- Requirement ID: `{req['requirementId'] or 'n/a'}`",
                    f"- Current statement: {req['statement'] or 'not answered'}",
                    f"- Excel row: {req['excelRow']}",
                    "",
                    "**Requirement**",
                    "",
                    req["requirement"],
                    "",
                    "**Current explanation / Begruendung**",
                    "",
                    req["explanation"] or "_None captured in workbook._",
                    "",
                    "**Answering notes**",
                    "",
                    *[f"- {note}" for note in requirement_guidance(req["requirement"])],
                    "",
                ]
            )
            if req["securityTest"] or req["reviewResult"] or req["reviewExplanation"]:
                lines.extend(
                    [
                        "**Review / test context captured from workbook**",
                        "",
                        f"- Security test: {req['securityTest'] or 'n/a'}",
                        f"- Review result: {req['reviewResult'] or 'n/a'}",
                        f"- Review explanation: {req['reviewExplanation'] or 'n/a'}",
                        "",
                    ]
                )
        sheet_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
        summary["sheets"].append({"title": sheet_title, "markdown": str(sheet_path), "requirementCount": len(reqs)})

    index_path = out_dir / f"{slug}_index.md"
    index_path.write_text("\n".join(index_lines).rstrip() + "\n", encoding="utf-8")
    summary_path = out_dir / f"{slug}_markdown_summary.json"
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    return {"index": str(index_path), "summary": str(summary_path), "sheets": summary["sheets"]}


def parse_requirements_html(html_path):
    soup = BeautifulSoup(Path(html_path).read_text(encoding="utf-8"), "lxml")
    requirements = []
    current_sheet = ""
    headings = []
    current = None

    def flush():
        nonlocal current
        if current:
            current["details"] = "\n\n".join(current.pop("detailsParts", [])).strip()
            current["motivation"] = "\n\n".join(current.pop("motivationParts", [])).strip()
            requirements.append(current)
            current = None

    def append_body(text):
        if not text:
            return
        before, marker, after = text.partition("Motivation:")
        if marker:
            if before.strip():
                append_requirement_or_details(before.strip())
            current["motivationParts"].append(f"Motivation: {after.strip()}")
            return
        append_requirement_or_details(text)

    def append_requirement_or_details(text):
        if not current["requirement"]:
            current["requirement"] = text
        else:
            current["detailsParts"].append(text)

    for element in soup.body.find_all(["h2", "h3", "h4", "p"], recursive=True):
        text = one_line(element.get_text(" ", strip=True))
        if not text:
            continue
        if element.name == "h2":
            flush()
            current_sheet = text
            headings = []
            continue
        if element.name in {"h3", "h4"}:
            flush()
            headings.append(text)
            if len(headings) > 4:
                headings = headings[-4:]
            continue
        req_match = re.match(r"Req\s+(\d+)\s*(.*)", text)
        if req_match:
            flush()
            current = {
                "sheet": current_sheet,
                "section": " / ".join(headings),
                "requirementNumber": req_match.group(1),
                "requirement": req_match.group(2).strip(),
                "details": "",
                "detailsParts": [],
                "motivation": "",
                "motivationParts": [],
                "id": "",
                "requirementId": "",
            }
            continue
        if not current:
            continue
        if text.startswith("ID:"):
            current["id"] = text.removeprefix("ID:").strip()
        elif text.startswith("Requirement-ID:"):
            current["requirementId"] = text.removeprefix("Requirement-ID:").strip()
        else:
            append_body(text)
    flush()
    return requirements


def render_markdown_from_html(html_path, out_dir, slug):
    out_dir.mkdir(parents=True, exist_ok=True)
    requirements = parse_requirements_html(html_path)
    by_sheet = {}
    for req in requirements:
        by_sheet.setdefault(req["sheet"], []).append(req)

    summary = {
        "source": str(html_path),
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "requirementCount": len(requirements),
        "sheets": [],
    }
    index_lines = [
        f"# {slug} SoC Requirements From HTML Export",
        "",
        f"Source HTML: `{html_path}`",
        "",
        ANSWER_GUIDE.strip(),
        "",
        "## Sheets",
        "",
    ]
    for index, (sheet, reqs) in enumerate(by_sheet.items(), start=1):
        sheet_path = out_dir / f"{slug}_{index:02d}_{safe_filename(sheet)}.md"
        index_lines.append(f"- [{sheet}]({sheet_path.name}) - {len(reqs)} requirements")
        lines = [
            f"# {sheet}",
            "",
            f"Requirement count: {len(reqs)}",
            "",
            ANSWER_GUIDE.strip(),
            "",
            "## Requirements",
            "",
        ]
        for req in reqs:
            lines.extend(
                [
                    f"### Req {req['requirementNumber']}",
                    "",
                    f"- Requirement-ID: `{req['requirementId'] or 'n/a'}`",
                    f"- Export ID: `{req['id'] or 'n/a'}`",
                    f"- Section: {req['section'] or 'n/a'}",
                    "",
                    "**Requirement**",
                    "",
                    req["requirement"] or "_No requirement title captured in the HTML export._",
                    "",
                    "**Details / motivation / references from HTML export**",
                    "",
                    req.get("details") or "_No additional details captured._",
                    "",
                    req["motivation"] or "_No motivation captured._",
                    "",
                    "**How to answer this question**",
                    "",
                    *[f"- {note}" for note in requirement_guidance(req["requirement"] + " " + req.get("details", "") + " " + req["motivation"])],
                    "",
                ]
            )
        sheet_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
        summary["sheets"].append({"title": sheet, "markdown": str(sheet_path), "requirementCount": len(reqs)})

    index_path = out_dir / f"{slug}_html_index.md"
    index_path.write_text("\n".join(index_lines).rstrip() + "\n", encoding="utf-8")
    summary_path = out_dir / f"{slug}_html_summary.json"
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    raw_json_path = out_dir / f"{slug}_html_requirements.json"
    raw_json_path.write_text(json.dumps(requirements, indent=2, ensure_ascii=False), encoding="utf-8")
    return {
        "index": str(index_path),
        "summary": str(summary_path),
        "requirementsJson": str(raw_json_path),
        "requirementCount": len(requirements),
        "sheets": summary["sheets"],
    }


def read_manifest(path):
    with Path(path).open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


PREPARE_SOC_SHEET_SCRIPT = r"""(() => new Promise((resolve) => {
  const clean = (node) => (node?.textContent || "").replace(/\s+/g, " ").trim();
  const clicked = [];
  const openOnce = () => {
    let count = 0;
    Array.from(document.querySelectorAll(".expandable-teaser.clickable")).forEach((teaser) => {
      const text = clean(teaser);
      if (!text || /ausblenden/i.test(text)) {
        return;
      }
      if (/SoC Details|Begr.ndung|Notizbereich/i.test(text)) {
        teaser.click();
        clicked.push(text.slice(0, 80));
        count += 1;
      }
    });
    return count;
  };
  const pass = (remaining) => {
    const count = openOnce();
    if (remaining <= 1 || count === 0) {
      setTimeout(() => resolve({clicked, clickedCount: clicked.length}), 750);
      return;
    }
    setTimeout(() => pass(remaining - 1), 750);
  };
  pass(5);
}))()"""


FIELD_INSPECTION_SCRIPT = r"""(() => {
  const cleanText = (node) => {
    if (!node) {
      return "";
    }
    const clone = node.cloneNode(true);
    clone.querySelectorAll("script, style").forEach((child) => child.remove());
    return (clone.textContent || "")
      .replace(/jQuery\(function\s*\(\)\s*\{[\s\S]*?\}\);/g, "")
      .replace(/[ \t]+\n/g, "\n")
      .replace(/\n{3,}/g, "\n\n")
      .trim();
  };
  const valueOf = (node) => node ? (node.value || "") : "";
  const radioGroups = Array.from(new Set(
    Array.from(document.querySelectorAll('input[type="radio"][name*="idComplianceRadioBtns"]'))
      .map((input) => input.name)
  ));
  const optionLabels = radioGroups.length
    ? Array.from(document.querySelectorAll(`input[name="${CSS.escape(radioGroups[0])}"]`)).map((input) => ({
        value: input.value,
        label: (input.closest("label")?.innerText || input.parentElement?.innerText || "").trim()
      }))
    : [];
  const requirements = Array.from(document.querySelectorAll('fieldset[class*="soc-statement"]')).map((fieldset) => {
    const reqMatch = (fieldset.innerText || "").match(/Anforderung\s+(\d+)/);
    const requirementIdText = Array.from(fieldset.querySelectorAll('[id*="compactMigrationId"], [id*="j_idt1123"]'))
      .map((node) => (node.innerText || node.textContent || "").trim())
      .find((text) => text.includes("Requirement-ID")) || "";
    const requirementTextNode = Array.from(fieldset.querySelectorAll('[id]'))
      .find((node) => /:j_idt1076:j_idt1078$/.test(node.id));
    const detailRootId = requirementTextNode?.id ? requirementTextNode.id.split(":").slice(0, -1).join(":") : "";
    const detailRoot = detailRootId ? document.getElementById(detailRootId) : null;
    const detailBody = detailRootId ? document.getElementById(`${detailRootId}:expandableBody`) : null;
    const detailText = cleanText(detailBody || detailRoot);
    const heading = reqMatch ? `Anforderung ${reqMatch[1]}` : "";
    let requirementText = (requirementTextNode?.innerText || requirementTextNode?.textContent || "").trim();
    if (!requirementText) {
      const withoutRequirementId = detailText.replace(/Requirement-ID:\s*[A-Za-z0-9_-]+/g, "").trim();
      const withoutHeading = heading ? withoutRequirementId.replace(new RegExp(`^${heading}\\s*`), "").trim() : withoutRequirementId;
      if (withoutHeading && withoutHeading !== heading) {
        requirementText = withoutHeading;
      }
    }
    const radios = Array.from(fieldset.querySelectorAll('input[type="radio"][name*="idComplianceRadioBtns"]'));
    const groupName = radios[0]?.name || "";
    const base = groupName.replace(/:idComplianceRadioBtns$/, "");
    const checkedStatement = radios.find((input) => input.checked);
    const checkedStatementLabel = checkedStatement
      ? (checkedStatement.closest("label")?.innerText || checkedStatement.parentElement?.innerText || "").trim()
      : "";
    const explanationTextarea = fieldset.querySelector('textarea[id*="idExplanationText"]');
    const noteTextarea = fieldset.querySelector('textarea[id*="idNoteInputField"]');
    const noteVisibleCheckbox = fieldset.querySelector('input[id*="idExternallyVisible_input"]');
    const externalNotesTextarea = fieldset.querySelector('textarea[id*="idExternText"]');
    return {
      heading,
      fieldsetId: fieldset.id,
      requirementId: (requirementIdText.match(/Requirement-ID:\s*([A-Za-z0-9_-]+)/) || [])[1] || "",
      requirementText,
      requirementDetailsText: detailText,
      currentStatementValue: checkedStatement?.value || "",
      currentStatementLabel: checkedStatementLabel,
      statementGroupName: groupName,
      statementOptions: radios.map((input) => ({
        id: input.id,
        value: input.value,
        checked: input.checked,
        disabled: input.disabled,
        label: (input.closest("label")?.innerText || input.parentElement?.innerText || "").trim()
      })),
      explanationTextareaId: explanationTextarea?.id || `${base}:idExplanationExpandable:idExplanationText`,
      explanationRendered: Boolean(explanationTextarea),
      explanationText: valueOf(explanationTextarea),
      noteTextareaId: noteTextarea?.id || "idNoteInputField in the per-requirement Notizbereich panel",
      noteRendered: Boolean(noteTextarea),
      noteText: valueOf(noteTextarea),
      noteVisibleCheckboxId: noteVisibleCheckbox?.id || "idExternallyVisible_input in the per-requirement Notizbereich panel",
      noteVisibleForOthers: noteVisibleCheckbox ? noteVisibleCheckbox.checked : null,
      visibleExternalNotesTextareaId: externalNotesTextarea?.id || "idExternText in the per-requirement Notizbereich panel",
      visibleExternalNotesText: valueOf(externalNotesTextarea),
      resetButtonId: fieldset.querySelector('button[id*="idResetComplianceButton"]')?.id || ""
    };
  });
  return {
    href: location.href,
    title: document.title,
    capturedAt: new Date().toISOString(),
    sheetTitle: (document.querySelector("h1,h2,.headline,.title")?.innerText || document.title || "").trim(),
    requirementCount: requirements.length,
    statementOptions: optionLabels,
    saveButtons: Array.from(document.querySelectorAll("button")).filter((button) => /Speichern/.test(button.innerText || ""))
      .map((button) => ({id: button.id, text: button.innerText.trim(), disabled: button.disabled})),
    filters: Array.from(document.querySelectorAll('input[type="checkbox"], select')).filter((el) =>
      (el.id || "").includes("Filter") || (el.id || "").includes("securityTestResultFilter")
    ).map((el) => ({
      tag: el.tagName,
      id: el.id,
      name: el.name,
      checked: el.checked,
      value: el.value,
      text: (el.innerText || el.textContent || "").trim()
    })),
    chapterNaButtons: Array.from(document.querySelectorAll('button[id*="idMakeChapterNAButton"]'))
      .map((button) => ({id: button.id, text: button.innerText.trim(), disabled: button.disabled})),
    requirements
  };
})()"""


def prepare_soc_sheet_for_capture(cdp):
    return cdp.evaluate_json(PREPARE_SOC_SHEET_SCRIPT, await_promise=True)


def inspect_soc_fields(cdp):
    return cdp.evaluate_json(FIELD_INSPECTION_SCRIPT)


def write_field_inspection_docs(inspections, out_dir, slug):
    out_dir.mkdir(parents=True, exist_ok=True)
    details_dir = out_dir / "sheets"
    details_dir.mkdir(parents=True, exist_ok=True)
    for index, inspection in enumerate(inspections, start=1):
        sheet_name = inspection.get("manifestName") or inspection.get("title") or f"sheet-{index}"
        detail_path = details_dir / f"{slug}_{index:02d}_{safe_filename(sheet_name)}.md"
        detail_lines = [
            f"# {sheet_name}",
            "",
            f"- URL: {inspection['href']}",
            f"- Requirements visible in browser: {inspection['requirementCount']}",
            f"- Sheet ID: `{inspection.get('sheet') or 'n/a'}`",
            f"- System ID: `{inspection.get('system') or 'n/a'}`",
            "",
            ANSWER_GUIDE.strip(),
            "",
            "## Requirements",
            "",
        ]
        for req in inspection.get("requirements", []):
            detail_lines.extend(
                [
                    f"### {req.get('heading') or 'Requirement'}",
                    "",
                    f"- Requirement ID: `{req.get('requirementId') or 'n/a'}`",
                    f"- Current statement: {req.get('currentStatementLabel') or req.get('currentStatementValue') or 'not answered'}",
                    f"- Statement field: `{req.get('statementGroupName')}`",
                    f"- Begruendung field: `{req.get('explanationTextareaId')}`",
                    f"- Note field: `{req.get('noteTextareaId')}`",
                    "",
                    "**Requirement**",
                    "",
                    md_escape(req.get("requirementText")) or "_No requirement text captured._",
                    "",
                    "**SoC details / motivation captured from browser**",
                    "",
                    md_escape(req.get("requirementDetailsText")) or "_No expanded details captured._",
                    "",
                    "**Captured Begruendung**",
                    "",
                    md_escape(req.get("explanationText")) or "_Empty or not rendered in the portal DOM._",
                    "",
                    "**Captured notes**",
                    "",
                    f"- Meine Notiz: {md_escape(req.get('noteText')) or '_Empty or not rendered in the portal DOM._'}",
                    f"- Fuer andere sichtbar: {req.get('noteVisibleForOthers') if req.get('noteVisibleForOthers') is not None else 'not rendered'}",
                    f"- Sichtbare Notizen anderer Nutzer: {md_escape(req.get('visibleExternalNotesText')) or '_Empty or not rendered in the portal DOM._'}",
                    "",
                    "**How to answer this question**",
                    "",
                    *[f"- {note}" for note in requirement_guidance(req.get("requirementText", ""))],
                    "",
                ]
            )
        detail_path.write_text("\n".join(detail_lines).rstrip() + "\n", encoding="utf-8")
        inspection["detailMarkdown"] = str(detail_path)

    json_path = out_dir / f"{slug}_field_inspection.json"
    json_path.write_text(json.dumps(inspections, indent=2, ensure_ascii=False), encoding="utf-8")
    lines = [
        f"# {slug} PSA SoC Browser Field Guide",
        "",
        "This guide was generated from the live PSA Portal DOM through Chrome DevTools. It captures expanded SoC details, current Begruendung values, notes, and the fields used to answer each requirement. It does not submit or save data.",
        "",
        ANSWER_GUIDE.strip(),
        "",
        "## Common Fields",
        "",
        "- Per requirement `Statement` radio group: `SOC`, `SOPC`, `SONC`, `NA`.",
        "- Per requirement `Begruendung` textarea: present as `idExplanationText`; required by process for `SOPC`, `SONC`, and `NA`.",
        "- Per requirement note panel: `idNoteInputField`, `idExternallyVisible_input`, and read-only `idExternText`.",
        "- Sheet actions: top/bottom `Speichern`; `Speichern und beenden` becomes available only when completion rules are met.",
        "- Chapter action: `Nicht zutreffend` can mark a chapter as not applicable, but the dialog requires a justification and should only be used intentionally.",
        "",
        "## Inspected Sheets",
        "",
    ]
    for inspection in inspections:
        lines.extend(
            [
                f"### {inspection.get('manifestName') or inspection.get('title')}",
                "",
                f"- URL: {inspection['href']}",
                f"- Detailed Markdown: [{Path(inspection['detailMarkdown']).name}](sheets/{Path(inspection['detailMarkdown']).name})",
                f"- Requirements visible in browser: {inspection['requirementCount']}",
                f"- Save buttons: {', '.join(button['text'] for button in inspection.get('saveButtons', [])) or 'none'}",
                f"- Chapter NA buttons: {len(inspection.get('chapterNaButtons', []))}",
                "",
                "| Requirement | Requirement ID | Current statement | Statement field | Begruendung captured | Note captured |",
                "| --- | --- | --- | --- | --- | --- |",
            ]
        )
        for req in inspection.get("requirements", []):
            lines.append(
                "| "
                + " | ".join(
                    [
                        one_line(req.get("heading") or "Requirement"),
                        f"`{one_line(req.get('requirementId')) or 'n/a'}`",
                        one_line(req.get("currentStatementLabel") or req.get("currentStatementValue") or "not answered"),
                        f"`{one_line(req.get('statementGroupName'))}`",
                        "yes" if md_escape(req.get("explanationText")) else "no",
                        "yes" if md_escape(req.get("noteText")) or md_escape(req.get("visibleExternalNotesText")) else "no",
                    ]
                )
                + " |"
            )
        lines.append("")
    md_path = out_dir / f"{slug}_field_guide.md"
    md_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return {"json": str(json_path), "markdown": str(md_path), "sheetCount": len(inspections)}


def download_socs(cdp, download_dir, timeout=120):
    download_dir.mkdir(parents=True, exist_ok=True)
    cdp.call(
        "Page.setDownloadBehavior",
        {"behavior": "allow", "downloadPath": str(download_dir.resolve())},
    )
    selected = select_all_socs(cdp)
    if not selected.get("ok"):
        return {"downloaded": False, "selection": selected}
    if selected["total"] == 0 or selected["checked"] == 0 or selected["exportDisabled"]:
        return {"downloaded": False, "selection": selected}

    start_time = time.time()
    click_result = click_bulk_export(cdp)
    if not click_result.get("ok"):
        return {"downloaded": False, "selection": selected, "click": click_result}

    downloaded = wait_for_download(download_dir, start_time, timeout=timeout)
    return {
        "downloaded": True,
        "selection": selected,
        "file": str(downloaded),
        "bytes": downloaded.stat().st_size,
    }


def command_snapshot(args):
    cdp = CdpClient(args.port)
    try:
        if args.system_id:
            navigate(cdp, system_url(args.base_url, args.system_id))
        wait_ready(cdp)
        snapshot = snapshot_page(cdp)
        result = write_snapshot(snapshot, Path(args.out_dir), args.slug)
        print(json.dumps(result, indent=2, ensure_ascii=False))
    finally:
        cdp.close()


def command_download(args):
    cdp = CdpClient(args.port)
    try:
        if args.system_id:
            navigate(cdp, system_url(args.base_url, args.system_id))
        wait_ready(cdp)
        snapshot = snapshot_page(cdp)
        snapshot_result = write_snapshot(snapshot, Path(args.snapshot_dir), args.slug)
        download_result = download_socs(cdp, Path(args.download_dir), timeout=args.timeout)
        print(json.dumps({"snapshot": snapshot_result, "download": download_result}, indent=2, ensure_ascii=False))
    finally:
        cdp.close()


def command_systems(args):
    cdp = CdpClient(args.port)
    try:
        navigate(cdp, f"{args.base_url.rstrip('/')}/my/systems/systems.xhtml")
        snapshot = snapshot_page(cdp)
        rows = []
        for link in snapshot["links"]:
            if "details/details.xhtml?system=" not in link["href"]:
                continue
            ids = parse_sheet_ids(link["href"])
            match = re.search(r"system=(\\d+)", link["href"])
            rows.append({"name": link["text"], "system": match.group(1) if match else "", "href": link["href"]})
        print(json.dumps(rows, indent=2, ensure_ascii=False))
    finally:
        cdp.close()


def command_extract_xlsx(args):
    result = extract_xlsx(Path(args.xlsx), Path(args.out_dir), args.slug)
    print(json.dumps(result, indent=2, ensure_ascii=False))


def command_render_markdown(args):
    result = render_markdown_from_xlsx(Path(args.xlsx), Path(args.out_dir), args.slug)
    print(json.dumps(result, indent=2, ensure_ascii=False))


def command_render_html_markdown(args):
    result = render_markdown_from_html(Path(args.html), Path(args.out_dir), args.slug)
    print(json.dumps(result, indent=2, ensure_ascii=False))


def command_inspect_fields(args):
    manifest = read_manifest(args.manifest)
    if args.limit:
        manifest = manifest[: args.limit]
    inspections = []
    cdp = CdpClient(args.port)
    try:
        for row in manifest:
            href = row.get("href")
            if not href:
                continue
            navigate(cdp, href)
            wait_ready(cdp, expected_url=href)
            prepare_result = prepare_soc_sheet_for_capture(cdp)
            inspection = inspect_soc_fields(cdp)
            inspection["prepareResult"] = prepare_result
            inspection["manifestName"] = row.get("name", "")
            inspection["sheet"] = row.get("sheet", "")
            inspection["system"] = row.get("system", "")
            inspections.append(inspection)
            time.sleep(args.pause)
        result = write_field_inspection_docs(inspections, Path(args.out_dir), args.slug)
        print(json.dumps(result, indent=2, ensure_ascii=False))
    finally:
        cdp.close()


def build_parser():
    parser = argparse.ArgumentParser(description="Capture PSA Portal system SoC snapshots and downloads via Chrome DevTools.")
    parser.add_argument("--port", type=int, default=9223, help="Chrome DevTools port.")
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL, help="PSA intranet UI base URL.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    snapshot = subparsers.add_parser("snapshot", help="Capture the current or specified system page.")
    snapshot.add_argument("--system-id", help="PSA system ID to navigate to before capture.")
    snapshot.add_argument("--slug", required=True, help="Output file prefix.")
    snapshot.add_argument("--out-dir", required=True, help="Snapshot output directory.")
    snapshot.set_defaults(func=command_snapshot)

    download = subparsers.add_parser("download-socs", help="Snapshot a system page and download all available SoC sheets as ZIP.")
    download.add_argument("--system-id", help="PSA system ID to navigate to before capture.")
    download.add_argument("--slug", required=True, help="Output file prefix.")
    download.add_argument("--snapshot-dir", required=True, help="Snapshot output directory.")
    download.add_argument("--download-dir", required=True, help="Download output directory.")
    download.add_argument("--timeout", type=int, default=120, help="Download timeout in seconds.")
    download.set_defaults(func=command_download)

    systems = subparsers.add_parser("list-systems", help="List system detail links from Meine Systeme.")
    systems.set_defaults(func=command_systems)

    extract = subparsers.add_parser("extract-xlsx", help="Extract a downloaded SoC XLSX into JSON and CSV files.")
    extract.add_argument("--xlsx", required=True, help="Downloaded SoC XLSX path.")
    extract.add_argument("--slug", required=True, help="Output file prefix.")
    extract.add_argument("--out-dir", required=True, help="Extraction output directory.")
    extract.set_defaults(func=command_extract_xlsx)

    render = subparsers.add_parser("render-markdown", help="Render SoC workbook requirements into Markdown docs.")
    render.add_argument("--xlsx", required=True, help="Downloaded SoC XLSX path.")
    render.add_argument("--slug", required=True, help="Output file prefix.")
    render.add_argument("--out-dir", required=True, help="Markdown output directory.")
    render.set_defaults(func=command_render_markdown)

    render_html = subparsers.add_parser("render-html-markdown", help="Render downloaded SoC HTML export into Markdown docs.")
    render_html.add_argument("--html", required=True, help="Downloaded SoC HTML export path.")
    render_html.add_argument("--slug", required=True, help="Output file prefix.")
    render_html.add_argument("--out-dir", required=True, help="Markdown output directory.")
    render_html.set_defaults(func=command_render_html_markdown)

    inspect = subparsers.add_parser("inspect-fields", help="Inspect live PSA SoC form fields from manifest links.")
    inspect.add_argument("--manifest", required=True, help="SoC manifest TSV with href column.")
    inspect.add_argument("--slug", required=True, help="Output file prefix.")
    inspect.add_argument("--out-dir", required=True, help="Field guide output directory.")
    inspect.add_argument("--limit", type=int, help="Only inspect the first N manifest rows.")
    inspect.add_argument("--pause", type=float, default=0.5, help="Pause between sheet inspections.")
    inspect.set_defaults(func=command_inspect_fields)
    return parser


def main():
    args = build_parser().parse_args()
    try:
        args.func(args)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
