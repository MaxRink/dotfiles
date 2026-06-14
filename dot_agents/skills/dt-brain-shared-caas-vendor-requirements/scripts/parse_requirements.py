#!/usr/bin/env python3
"""
Parse vendor infrastructure requirement Excel files into standardized JSON.

Supports:
  - Ericsson SDM/CCDM format (multi-sheet, Category/Description/Requirement/Consequence columns)
  - Generic format (auto-detects requirement-like columns)

Usage:
  parse_requirements.py <excel-file> [--output FILE] [--format json|table]
  parse_requirements.py <excel-file> --category-col A --requirement-col C --consequence-col D
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


# Known header patterns for auto-detection
REQUIREMENT_HEADERS = [
    'requirement', 'requirements', 'req', 'description', 'requirement description',
    'requirement text', 'spec', 'specification'
]
CATEGORY_HEADERS = [
    'category', 'area', 'domain', 'section', 'type', 'group'
]
CONSEQUENCE_HEADERS = [
    'consequence', 'impact', 'risk', 'effect', 'result', 'consequences'
]
COMMENT_HEADERS = [
    'comment', 'comments', 'note', 'notes', 'remark', 'remarks',
    'customer', 'response', 'feedback'
]
# Sheets to skip
SKIP_SHEETS = ['summary', 'overview', 'toc', 'table of contents', 'cover', 'index', 'changelog']


def detect_columns(ws):
    """Auto-detect column roles from header row."""
    headers = {}
    header_row = None

    # Check first 5 rows for headers
    for row_idx in range(1, 6):
        row_values = []
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                row_values.append((col_idx, str(val).strip().lower()))

        # If at least 2 cells look like headers, this is probably the header row
        header_like = sum(1 for _, v in row_values if any(h in v for h in
            REQUIREMENT_HEADERS + CATEGORY_HEADERS + CONSEQUENCE_HEADERS + COMMENT_HEADERS))
        if header_like >= 2:
            header_row = row_idx
            for col_idx, val in row_values:
                if any(h in val for h in REQUIREMENT_HEADERS) and 'requirement_col' not in headers:
                    headers['requirement_col'] = col_idx
                elif any(h in val for h in CATEGORY_HEADERS) and 'category_col' not in headers:
                    headers['category_col'] = col_idx
                elif any(h in val for h in CONSEQUENCE_HEADERS) and 'consequence_col' not in headers:
                    headers['consequence_col'] = col_idx
                elif any(h in val for h in COMMENT_HEADERS) and 'comment_col' not in headers:
                    headers['comment_col'] = col_idx
            break

    return headers, header_row


def is_ericsson_sdm_format(wb):
    """Check if workbook matches Ericsson SDM/CCDM format."""
    sheet_names_lower = [s.lower() for s in wb.sheetnames]
    sdm_indicators = ['runtime environment', 'compute requirements', 'network requirements',
                      'storage requirements', 'security requirements', 'recommendations']
    matches = sum(1 for ind in sdm_indicators if any(ind in s for s in sheet_names_lower))
    return matches >= 3


def parse_ericsson_sdm(wb, verbose=False):
    """Parse Ericsson SDM/CCDM format workbook."""
    requirements = []
    req_id = 0

    for sheet_name in wb.sheetnames:
        if sheet_name.lower() in SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        if ws.max_row is None or ws.max_row < 3:
            continue

        # Ericsson SDM format: Row 1 = headers, data starts row 2
        # Columns: A=Category, B=Description (sub-items), C=Requirement, D=Consequence, E+=Comments
        # Special: Runtime Environment sheet has parent items in A with sub-items in B

        # Detect column layout
        detected, header_row = detect_columns(ws)
        if not detected.get('requirement_col'):
            # Try Ericsson default: C=Requirement
            if ws.cell(row=1, column=3).value and 'requirement' in str(ws.cell(row=1, column=3).value).lower():
                detected = {'category_col': 1, 'requirement_col': 3, 'consequence_col': 4}
                header_row = 1
            else:
                if verbose:
                    print(f"  Skipping sheet '{sheet_name}': no requirement column detected", file=sys.stderr)
                continue

        cat_col = detected.get('category_col', 1)
        req_col = detected.get('requirement_col', 3)
        cons_col = detected.get('consequence_col', 4)
        comment_col = detected.get('comment_col')

        # Also check for T-CaaS column (typically last meaningful column)
        tcaas_col = None
        if header_row:
            for col_idx in range(1, ws.max_column + 1):
                val = ws.cell(row=header_row, column=col_idx).value
                if val and 'caas' in str(val).lower():
                    tcaas_col = col_idx
                    break

        # Find customer comment column (often has NatCo name)
        if not comment_col and header_row:
            for col_idx in range(1, ws.max_column + 1):
                val = ws.cell(row=header_row, column=col_idx).value
                if val and col_idx not in [cat_col, req_col, cons_col, tcaas_col]:
                    val_lower = str(val).strip().lower()
                    if any(h in val_lower for h in COMMENT_HEADERS) or 'comment' in val_lower:
                        comment_col = col_idx
                        break

        data_start = (header_row or 1) + 1

        # Track parent items for consolidation
        current_parent = None
        parent_subitems = []

        for row_idx in range(data_start, (ws.max_row or 0) + 1):
            cat_val = ws.cell(row=row_idx, column=cat_col).value
            req_val = ws.cell(row=row_idx, column=req_col).value
            cons_val = ws.cell(row=row_idx, column=cons_col).value if cons_col else None
            comment_val = ws.cell(row=row_idx, column=comment_col).value if comment_col else None

            # Check for sub-items in column B (description column in Ericsson format)
            desc_val = None
            if cat_col == 1 and req_col >= 3:
                desc_val = ws.cell(row=row_idx, column=2).value

            # Skip completely empty rows
            if not any([cat_val, req_val, desc_val]):
                # If we have accumulated sub-items, flush the parent
                if current_parent and parent_subitems:
                    req_id += 1
                    parent_text = current_parent['text']
                    # Consolidate sub-items into parent
                    sub_text = '\n- '.join(parent_subitems)
                    full_text = f"{parent_text}\n- {sub_text}"
                    requirements.append({
                        'id': req_id,
                        'sheet': sheet_name,
                        'row_num': current_parent['row'],
                        'category': current_parent.get('category', ''),
                        'requirement_text': full_text.strip(),
                        'consequence': current_parent.get('consequence', ''),
                        'customer_comment': current_parent.get('comment', ''),
                    })
                    current_parent = None
                    parent_subitems = []
                continue

            # Sub-item row: has description but no category or requirement
            if desc_val and not cat_val and not req_val:
                if current_parent:
                    parent_subitems.append(str(desc_val).strip())
                continue

            # If we had a parent accumulating sub-items, flush it
            if current_parent and parent_subitems:
                req_id += 1
                parent_text = current_parent['text']
                sub_text = '\n- '.join(parent_subitems)
                full_text = f"{parent_text}\n- {sub_text}"
                requirements.append({
                    'id': req_id,
                    'sheet': sheet_name,
                    'row_num': current_parent['row'],
                    'category': current_parent.get('category', ''),
                    'requirement_text': full_text.strip(),
                    'consequence': current_parent.get('consequence', ''),
                    'customer_comment': current_parent.get('comment', ''),
                })
                current_parent = None
                parent_subitems = []

            # Regular requirement row
            req_text = str(req_val).strip() if req_val else ''
            cat_text = str(cat_val).strip() if cat_val else ''
            cons_text = str(cons_val).strip() if cons_val else ''
            comment_text = str(comment_val).strip() if comment_val else ''

            if not req_text and not cat_text:
                continue

            # Check if this might be a parent row (has category but requirement
            # might have sub-items following)
            # Look ahead to see if next row has desc_val but no cat/req
            if cat_col == 1 and req_col >= 3:
                next_row = row_idx + 1
                if next_row <= (ws.max_row or 0):
                    next_desc = ws.cell(row=next_row, column=2).value
                    next_cat = ws.cell(row=next_row, column=cat_col).value
                    next_req = ws.cell(row=next_row, column=req_col).value
                    if next_desc and not next_cat and not next_req:
                        # This is a parent row - start accumulating
                        current_parent = {
                            'text': req_text or cat_text,
                            'row': row_idx,
                            'category': cat_text,
                            'consequence': cons_text,
                            'comment': comment_text,
                        }
                        continue

            # Standalone requirement
            req_id += 1
            requirements.append({
                'id': req_id,
                'sheet': sheet_name,
                'row_num': row_idx,
                'category': cat_text,
                'requirement_text': req_text if req_text else cat_text,
                'consequence': cons_text,
                'customer_comment': comment_text,
            })

        # Flush any remaining parent
        if current_parent:
            req_id += 1
            if parent_subitems:
                sub_text = '\n- '.join(parent_subitems)
                full_text = f"{current_parent['text']}\n- {sub_text}"
            else:
                full_text = current_parent['text']
            requirements.append({
                'id': req_id,
                'sheet': sheet_name,
                'row_num': current_parent['row'],
                'category': current_parent.get('category', ''),
                'requirement_text': full_text.strip(),
                'consequence': current_parent.get('consequence', ''),
                'customer_comment': current_parent.get('comment', ''),
            })

    return requirements


def parse_generic(wb, col_overrides=None, verbose=False):
    """Parse generic Excel format with auto-detected or overridden columns."""
    requirements = []
    req_id = 0

    for sheet_name in wb.sheetnames:
        if sheet_name.lower() in SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        if ws.max_row is None or ws.max_row < 2:
            continue

        if col_overrides:
            detected = col_overrides
            header_row = 1
        else:
            detected, header_row = detect_columns(ws)

        if not detected.get('requirement_col'):
            if verbose:
                print(f"  Skipping sheet '{sheet_name}': no requirement column detected", file=sys.stderr)
            continue

        req_col = detected['requirement_col']
        cat_col = detected.get('category_col')
        cons_col = detected.get('consequence_col')
        comment_col = detected.get('comment_col')

        data_start = (header_row or 1) + 1

        for row_idx in range(data_start, (ws.max_row or 0) + 1):
            req_val = ws.cell(row=row_idx, column=req_col).value
            if not req_val:
                continue

            req_id += 1
            requirements.append({
                'id': req_id,
                'sheet': sheet_name,
                'row_num': row_idx,
                'category': str(ws.cell(row=row_idx, column=cat_col).value or '').strip() if cat_col else '',
                'requirement_text': str(req_val).strip(),
                'consequence': str(ws.cell(row=row_idx, column=cons_col).value or '').strip() if cons_col else '',
                'customer_comment': str(ws.cell(row=row_idx, column=comment_col).value or '').strip() if comment_col else '',
            })

    return requirements


def col_letter_to_num(letter):
    """Convert column letter (A, B, ..., Z, AA, ...) to 1-based number."""
    result = 0
    for ch in letter.upper():
        result = result * 26 + (ord(ch) - ord('A') + 1)
    return result


def main():
    parser = argparse.ArgumentParser(
        description='Parse vendor infrastructure requirements from Excel into JSON',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  %(prog)s requirements.xlsx
  %(prog)s requirements.xlsx --output /tmp/reqs.json
  %(prog)s requirements.xlsx --category-col A --requirement-col C
  %(prog)s requirements.xlsx --format table
        """
    )

    parser.add_argument('excel_file', help='Path to the Excel file')
    parser.add_argument('--output', '-o', help='Output file path (default: stdout)')
    parser.add_argument('--format', choices=['json', 'table'], default='json',
                        help='Output format (default: json)')
    parser.add_argument('--category-col', help='Column letter for category (e.g., A)')
    parser.add_argument('--requirement-col', help='Column letter for requirement text (e.g., C)')
    parser.add_argument('--consequence-col', help='Column letter for consequence (e.g., D)')
    parser.add_argument('--comment-col', help='Column letter for customer comment (e.g., E)')
    parser.add_argument('--verbose', '-v', action='store_true', help='Verbose output')

    args = parser.parse_args()

    excel_path = Path(args.excel_file)
    if not excel_path.exists():
        print(f"Error: File not found: {excel_path}", file=sys.stderr)
        sys.exit(1)

    if args.verbose:
        print(f"Loading: {excel_path}", file=sys.stderr)

    wb = load_workbook(str(excel_path), data_only=True, read_only=True)

    if args.verbose:
        print(f"Sheets: {', '.join(wb.sheetnames)}", file=sys.stderr)

    # Build column overrides if specified
    col_overrides = None
    if args.requirement_col:
        col_overrides = {
            'requirement_col': col_letter_to_num(args.requirement_col),
        }
        if args.category_col:
            col_overrides['category_col'] = col_letter_to_num(args.category_col)
        if args.consequence_col:
            col_overrides['consequence_col'] = col_letter_to_num(args.consequence_col)
        if args.comment_col:
            col_overrides['comment_col'] = col_letter_to_num(args.comment_col)

    # Detect format and parse
    if col_overrides:
        requirements = parse_generic(wb, col_overrides, args.verbose)
    elif is_ericsson_sdm_format(wb):
        if args.verbose:
            print("Detected: Ericsson SDM/CCDM format", file=sys.stderr)
        requirements = parse_ericsson_sdm(wb, args.verbose)
    else:
        if args.verbose:
            print("Using: Generic format with auto-detection", file=sys.stderr)
        requirements = parse_generic(wb, verbose=args.verbose)

    wb.close()

    if args.verbose:
        print(f"Parsed: {len(requirements)} requirements", file=sys.stderr)
        sheets = {}
        for r in requirements:
            sheets[r['sheet']] = sheets.get(r['sheet'], 0) + 1
        for sheet, count in sheets.items():
            print(f"  {sheet}: {count}", file=sys.stderr)

    # Output
    if args.format == 'table':
        # Print as table
        print(f"{'ID':>4}  {'Sheet':<35}  {'Category':<25}  {'Requirement':<60}  {'Comment':<30}")
        print('-' * 160)
        for r in requirements:
            req_short = r['requirement_text'][:57] + '...' if len(r['requirement_text']) > 60 else r['requirement_text']
            comment_short = r['customer_comment'][:27] + '...' if len(r['customer_comment']) > 30 else r['customer_comment']
            print(f"{r['id']:>4}  {r['sheet']:<35}  {r['category']:<25}  {req_short:<60}  {comment_short:<30}")
    else:
        output = json.dumps(requirements, indent=2, ensure_ascii=False)
        if args.output:
            Path(args.output).write_text(output)
            print(f"Saved {len(requirements)} requirements to {args.output}", file=sys.stderr)
        else:
            print(output)


if __name__ == '__main__':
    main()
