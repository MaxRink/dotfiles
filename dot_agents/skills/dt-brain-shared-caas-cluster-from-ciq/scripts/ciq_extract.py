#!/usr/bin/env python3
"""
CiQ Excel Extractor — Helper for cluster-from-ciq skill.

Lists sheets, dumps key rows from each sheet, and attempts to identify
the vendor from sheet/cell content.

Usage:
    ciq_extract.py <excel_file> [--sheet <name>] [--rows <max_rows>]

Examples:
    ciq_extract.py CiQ-file.xlsx                     # Overview of all sheets
    ciq_extract.py CiQ-file.xlsx --sheet "IP Subnet"  # Dump specific sheet
    ciq_extract.py CiQ-file.xlsx --rows 80            # More rows per sheet
"""

import argparse
import sys

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


VENDOR_KEYWORDS = {
    "mavenir": ["mavenir", "mavenir_core", "upf-dpe", "amf-comm", "smf-plc", "isupf", "isme", "isepdg"],
    "nokia": ["nokia", "nokia_core", "sdv", "cbis", "ncs", "netact"],
    "ericsson": ["ericsson", "ericsson_core", "eccd", "ccd", "enm"],
}


def detect_vendor(wb):
    """Scan first 20 rows of each sheet for vendor keywords."""
    scores = {v: 0 for v in VENDOR_KEYWORDS}
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
            for cell in row:
                if cell is None:
                    continue
                text = str(cell).lower()
                for vendor, keywords in VENDOR_KEYWORDS.items():
                    for kw in keywords:
                        if kw in text:
                            scores[vendor] += 1
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else "unknown"


def list_sheets(wb):
    """Print sheet names with dimensions."""
    print("=" * 60)
    print("SHEETS OVERVIEW")
    print("=" * 60)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"  {name:<45} ({ws.max_row}x{ws.max_column})")
    print()


def dump_sheet(ws, sheet_name, max_rows=50):
    """Dump non-empty rows from a sheet."""
    print(f"\n{'=' * 60}")
    print(f"SHEET: {sheet_name} ({ws.max_row}x{ws.max_column})")
    print(f"{'=' * 60}")
    row_count = 0
    for row in ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row), values_only=False):
        vals = [(c.coordinate, c.value) for c in row if c.value is not None]
        if vals:
            parts = [f"{coord}={repr(val)}" for coord, val in vals]
            print(f"  {' | '.join(parts)}")
            row_count += 1
    if row_count == 0:
        print("  (empty)")


# Sheets that typically contain CiQ data relevant for cluster definitions
PRIORITY_SHEETS = [
    "overview", "index", "platform_pre-requisites", "platform_prerequisites",
    "sriov", "hugepage", "ip subnet", "ip_subnet", "cnf_networks",
    "cit_cloud", "server", "distribution", "network config", "network_config",
    "vbom", "session config", "worker", "interfaces",
]


def is_priority_sheet(name):
    """Check if sheet name matches known priority patterns."""
    lower = name.lower()
    return any(p in lower for p in PRIORITY_SHEETS)


def main():
    parser = argparse.ArgumentParser(description="Extract CiQ data from Excel for cluster definition")
    parser.add_argument("excel_file", help="Path to CiQ Excel file")
    parser.add_argument("--sheet", help="Dump a specific sheet by name (partial match)")
    parser.add_argument("--rows", type=int, default=50, help="Max rows to dump per sheet (default: 50)")
    parser.add_argument("--all", action="store_true", help="Dump all sheets (not just priority ones)")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.excel_file, data_only=True)

    # Vendor detection
    vendor = detect_vendor(wb)
    print(f"Detected vendor: {vendor.upper()}")
    print()

    # Sheet listing
    list_sheets(wb)

    if args.sheet:
        # Dump specific sheet (partial match)
        matches = [s for s in wb.sheetnames if args.sheet.lower() in s.lower()]
        if not matches:
            print(f"No sheet matching '{args.sheet}' found.")
            sys.exit(1)
        for name in matches:
            dump_sheet(wb[name], name, args.rows)
    elif args.all:
        for name in wb.sheetnames:
            dump_sheet(wb[name], name, args.rows)
    else:
        # Dump priority sheets only
        print("Dumping priority sheets (use --all for everything, --sheet <name> for specific):\n")
        dumped = []
        for name in wb.sheetnames:
            if is_priority_sheet(name):
                dump_sheet(wb[name], name, args.rows)
                dumped.append(name)
        if not dumped:
            print("No priority sheets detected. Use --all to dump everything.")


if __name__ == "__main__":
    main()
